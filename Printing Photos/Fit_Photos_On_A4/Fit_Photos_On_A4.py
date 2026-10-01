from pathlib import Path
from PIL import Image, ImageOps
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_TABLE_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import tempfile
import uuid

PHOTO_FOLDER = Path(
    r"path/to/photos"
)

OUTPUT_FILE = PHOTO_FOLDER / "photo_layout.docx"

PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7

MARGIN_CM = 1.0

GAP_CM = 1.0

BOTTOM_SAFETY_CM = 1.0

SIX_INCHES_CM = 15.24
FOUR_INCHES_CM = 10.16

SUPPORTED_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
}

USABLE_WIDTH_CM = PAGE_WIDTH_CM - (2 * MARGIN_CM)

USABLE_HEIGHT_CM = (
    PAGE_HEIGHT_CM
    - (2 * MARGIN_CM)
    - BOTTOM_SAFETY_CM
)

def get_image_info(path):

    with Image.open(path) as img:
        img = ImageOps.exif_transpose(img)

        width, height = img.size

    ratio = width / height

    return {
        "path": path,
        "pixel_width": width,
        "pixel_height": height,
        "ratio": ratio,
    }

def get_preferred_size(info):

    ratio = info["ratio"]

    if ratio >= 1:

        width = SIX_INCHES_CM
        height = width / ratio

        if height > FOUR_INCHES_CM:
            height = FOUR_INCHES_CM
            width = height * ratio

    else:

        height = SIX_INCHES_CM
        width = height * ratio

        if width > FOUR_INCHES_CM:
            width = FOUR_INCHES_CM
            height = width / ratio

    return width, height

def create_word_safe_copy(path, temp_folder):

    with Image.open(path) as img:

        img = ImageOps.exif_transpose(img)

        if img.mode == "RGBA":
            background = Image.new(
                "RGB",
                img.size,
                "white"
            )

            background.paste(
                img,
                mask=img.getchannel("A")
            )

            img = background

        elif img.mode != "RGB":
            img = img.convert("RGB")

        unique_name = (
            f"{path.stem}_"
            f"{uuid.uuid4().hex[:8]}.jpg"
        )

        output_path = temp_folder / unique_name

        img.save(
            output_path,
            "JPEG",
            quality=95
        )

        return output_path

def fit_inside(
    width,
    height,
    max_width,
    max_height
):

    scale = min(
        1.0,
        max_width / width,
        max_height / height
    )

    return (
        width * scale,
        height * scale
    )

def make_grid_layout(
    photos,
    rows,
    cols
):

    if rows * cols < len(photos):
        return None

    available_width = (
        USABLE_WIDTH_CM
        - GAP_CM * (cols - 1)
    )

    available_height = (
        USABLE_HEIGHT_CM
        - GAP_CM * (rows - 1)
    )

    if available_width <= 0:
        return None

    if available_height <= 0:
        return None

    cell_width = available_width / cols
    cell_height = available_height / rows

    placements = []

    total_area = 0

    for photo in photos:

        width, height = fit_inside(
            photo["preferred_width"],
            photo["preferred_height"],
            cell_width,
            cell_height
        )

        placements.append(
            {
                "photo": photo,
                "width": width,
                "height": height,
            }
        )

        total_area += width * height

    return {
        "type": "grid",
        "rows": rows,
        "cols": cols,
        "placements": placements,
        "cell_width": cell_width,
        "cell_height": cell_height,
        "score": total_area,
    }

def find_best_layout(photos):

    count = len(photos)

    possible_grids = {
        1: [
            (1, 1),
        ],

        2: [
            (1, 2),
            (2, 1),
        ],

        3: [
            (1, 3),
            (3, 1),
            (2, 2),
        ],

        4: [
            (2, 2),
            (1, 4),
            (4, 1),
        ],
    }

    layouts = []

    for rows, cols in possible_grids[count]:

        result = make_grid_layout(
            photos,
            rows,
            cols
        )

        if result is not None:
            layouts.append(result)

    return max(
        layouts,
        key=lambda layout: layout["score"]
    )

def choose_photos_for_page(
    photos,
    start_index
):

    remaining = len(photos) - start_index

    maximum = min(4, remaining)

    candidates = []

    for count in range(
        1,
        maximum + 1
    ):

        group = photos[
            start_index:
            start_index + count
        ]

        layout = find_best_layout(group)

        displayed_area = layout["score"]

        preferred_area = sum(
            photo["preferred_width"]
            * photo["preferred_height"]
            for photo in group
        )

        size_retention = (
            displayed_area / preferred_area
            if preferred_area
            else 0
        )

        count_bonus = {
            1: 1.00,
            2: 1.30,
            3: 1.40,
            4: 1.45,
        }[count]

        score = (
            displayed_area
            * size_retention
            * count_bonus
        )

        candidates.append(
            {
                "count": count,
                "layout": layout,
                "score": score,
                "size_retention": size_retention,
            }
        )

    best = max(
        candidates,
        key=lambda item: item["score"]
    )

    return (
        best["count"],
        best["layout"]
    )

def remove_table_borders(table):

    tbl = table._tbl

    tbl_pr = tbl.tblPr

    borders = tbl_pr.first_child_found_in(
        "w:tblBorders"
    )

    if borders is None:
        borders = OxmlElement(
            "w:tblBorders"
        )

        tbl_pr.append(borders)

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):

        element = borders.find(
            qn(f"w:{edge}")
        )

        if element is None:
            element = OxmlElement(
                f"w:{edge}"
            )

            borders.append(element)

        element.set(
            qn("w:val"),
            "nil"
        )

def set_cell_margins(
    cell,
    top=0,
    start=0,
    bottom=0,
    end=0
):

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in(
        "w:tcMar"
    )

    if tc_mar is None:
        tc_mar = OxmlElement(
            "w:tcMar"
        )

        tc_pr.append(tc_mar)

    values = {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }

    for name, value in values.items():

        node = tc_mar.find(
            qn(f"w:{name}")
        )

        if node is None:
            node = OxmlElement(
                f"w:{name}"
            )

            tc_mar.append(node)

        node.set(
            qn("w:w"),
            str(int(value))
        )

        node.set(
            qn("w:type"),
            "dxa"
        )

def prevent_row_split(row):

    tr_pr = row._tr.get_or_add_trPr()

    cant_split = OxmlElement(
        "w:cantSplit"
    )

    tr_pr.append(cant_split)

def remove_cell_paragraph_spacing(cell):

    for paragraph in cell.paragraphs:

        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0

        paragraph.paragraph_format.line_spacing = 1

def add_photo_to_cell(
    cell,
    image_path,
    width_cm,
    height_cm
):

    cell.text = ""

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )

    paragraph = cell.paragraphs[0]

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0

    run = paragraph.add_run()

    run.add_picture(
        str(image_path),
        width=Cm(width_cm),
        height=Cm(height_cm)
    )

def add_page_break_before_next_layout(
    document
):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0

    run = paragraph.add_run()

    run.add_break()

def add_layout_to_document(
    document,
    layout,
    image_lookup
):

    rows = layout["rows"]
    cols = layout["cols"]

    table = document.add_table(
        rows=rows,
        cols=cols
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.autofit = False

    remove_table_borders(table)

    cell_width = layout["cell_width"]
    cell_height = layout["cell_height"]

    for row in table.rows:

        prevent_row_split(row)

        row.height = Cm(cell_height)

        row.height_rule = (
            WD_ROW_HEIGHT_RULE.EXACTLY
        )

    for row in table.rows:
        for cell in row.cells:

            cell.width = Cm(cell_width)

            set_cell_margins(
                cell,
                top=0,
                bottom=0,
                start=0,
                end=0
            )

            remove_cell_paragraph_spacing(cell)

    for index, placement in enumerate(
        layout["placements"]
    ):

        row_index = index // cols
        col_index = index % cols

        cell = table.cell(
            row_index,
            col_index
        )

        photo = placement["photo"]

        add_photo_to_cell(
            cell,
            image_lookup[photo["path"]],
            placement["width"],
            placement["height"]
        )

    trailing_paragraph = (
        document.add_paragraph()
    )

    trailing_paragraph.paragraph_format.space_before = 0
    trailing_paragraph.paragraph_format.space_after = 0
    trailing_paragraph.paragraph_format.line_spacing = 1

def main():

    print()
    print("Photo layout generator")
    print("======================")
    print()

    if not PHOTO_FOLDER.exists():

        raise FileNotFoundError(
            "\nPhoto folder does not exist:\n"
            f"{PHOTO_FOLDER}"
        )

    photo_paths = sorted(
        [
            path
            for path in PHOTO_FOLDER.iterdir()
            if (
                path.is_file()
                and
                path.suffix.lower()
                in SUPPORTED_EXTENSIONS
            )
        ],
        key=lambda path: path.name.lower()
    )

    if not photo_paths:

        raise RuntimeError(
            "\nNo supported photographs were found in:\n"
            f"{PHOTO_FOLDER}"
        )

    print(
        f"Found {len(photo_paths)} photographs."
    )

    print()

    photos = []

    for path in photo_paths:

        try:

            info = get_image_info(path)

            preferred_width, preferred_height = (
                get_preferred_size(info)
            )

            info["preferred_width"] = (
                preferred_width
            )

            info["preferred_height"] = (
                preferred_height
            )

            photos.append(info)

            print(
                f"{path.name}"
                f"  ->  "
                f"{preferred_width:.1f} cm"
                f" x "
                f"{preferred_height:.1f} cm"
            )

        except Exception as error:

            print()
            print(
                f"Skipping {path.name}"
            )

            print(
                f"Reason: {error}"
            )

            print()

    if not photos:

        raise RuntimeError(
            "None of the photographs could be read."
        )

    document = Document()

    section = document.sections[0]

    section.page_width = (
        Cm(PAGE_WIDTH_CM)
    )

    section.page_height = (
        Cm(PAGE_HEIGHT_CM)
    )

    section.top_margin = (
        Cm(MARGIN_CM)
    )

    section.bottom_margin = (
        Cm(MARGIN_CM)
    )

    section.left_margin = (
        Cm(MARGIN_CM)
    )

    section.right_margin = (
        Cm(MARGIN_CM)
    )

    section.header_distance = Cm(0.3)
    section.footer_distance = Cm(0.3)

    with tempfile.TemporaryDirectory() as temp_dir:

        temp_folder = Path(temp_dir)

        image_lookup = {}

        print()
        print(
            "Preparing photographs..."
        )
        print()

        for number, photo in enumerate(
            photos,
            start=1
        ):

            print(
                f"{number}/{len(photos)} "
                f"{photo['path'].name}"
            )

            safe_image = (
                create_word_safe_copy(
                    photo["path"],
                    temp_folder
                )
            )

            image_lookup[
                photo["path"]
            ] = safe_image

        print()
        print(
            "Creating document pages..."
        )
        print()

        current_index = 0

        page_number = 1

        while current_index < len(photos):

            count, layout = (
                choose_photos_for_page(
                    photos,
                    current_index
                )
            )

            if page_number > 1:

                add_page_break_before_next_layout(
                    document
                )

            print(
                f"Page {page_number}: "
                f"{count} photograph"
                f"{'s' if count != 1 else ''}"
                f"  |  "
                f"{layout['rows']} x "
                f"{layout['cols']} layout"
            )

            add_layout_to_document(
                document,
                layout,
                image_lookup
            )

            current_index += count
            page_number += 1

        document.save(
            OUTPUT_FILE
        )

    print()
    print("======================")
    print("Finished")
    print("======================")
    print()

    print(
        "Word document created:"
    )

    print(
        OUTPUT_FILE
    )

    print()

if __name__ == "__main__":
    main()
